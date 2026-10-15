# /// script
# requires-python = ">=3.11"
# dependencies = ["python-pptx>=1.0"]
# ///
"""Check that a PDF export carries the same on-slide text as the deck, slide by slide.

Needs `pdftotext` (poppler). A slide fails when more than 3% (minimum 3) of its words of
three or more characters are missing from the matching PDF page.

    uv run scripts/check_pdf_text.py slides/deck.pdf slides/deck.pptx
"""
import re,subprocess,sys
from pptx import Presentation
pdf,pptx=sys.argv[1],sys.argv[2]
norm=lambda s: re.sub(r"[^0-9a-z]+","",s.lower())
pres=Presentation(pptx); bad=0
for i,slide in enumerate(pres.slides,1):
    page=subprocess.run(["pdftotext","-f",str(i),"-l",str(i),pdf,"-"],capture_output=True,text=True).stdout
    P=norm(page); words=[]
    for sh in slide.shapes:
        if sh.has_text_frame: words+=sh.text_frame.text.split()
        if sh.has_table: words+=[w for r in sh.table.rows for c in r.cells for w in c.text.split()]
    toks=[norm(w) for w in words if len(norm(w))>=3]
    miss=[t for t in toks if t not in P]
    if len(miss)>max(2,0.03*len(toks)): bad+=1; print(i,len(toks),miss[:8])
print("slides",len(pres.slides),"with mismatches",bad); sys.exit(1 if bad else 0)
