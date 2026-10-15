# /// script
# requires-python = ">=3.11"
# dependencies = ["python-pptx>=1.0"]
# ///
"""Check that every mapped slide value matches its experiment result file.

For each entry in slides/values_map.json the value is recomputed from its source file
(same formats as slides/values.js) and must appear in the text of the named slide in the
built deck. Exits non-zero on any mismatch. Run from the repository root:

    uv run scripts/check_slide_values.py [slides/Knowledge_Graphs_Skills_Neuroimaging_v4.pptx]
"""

import json
import math
import sys
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
MINUS = "−"


def get(obj, dotted):
    if dotted in obj:
        return obj[dotted]
    for key in dotted.split("."):
        obj = obj[key]
    return obj


def half_up(x):
    """Math.round semantics: ties go towards +infinity."""
    return math.floor(x + 0.5)


FORMATS = {
    "int": lambda x: f"{half_up(x):,}",
    "round": lambda x: str(half_up(x)),
    "dec2": lambda x: f"{x:.2f}",
    "dec3": lambda x: f"{x:.3f}",
    "pct0": lambda x: f"{x * 100:.0f}",
    "pct1": lambda x: f"{x * 100:.1f}",
    "signed2": lambda x: (MINUS if x < 0 else "") + f"{abs(x):.2f}",
}


def expected(entry, cache):
    src = entry["source"]
    if src not in cache:
        cache[src] = json.loads((ROOT / src).read_text())
    data = cache[src]
    if entry["format"] == "sub":
        return FORMATS["int"](get(data, entry["path"]) - get(data, entry["path2"]))
    return FORMATS[entry["format"]](get(data, entry["path"]))


def slide_texts(pptx_path):
    """Return the on-slide text (not notes) of every slide."""
    texts = []
    for slide in Presentation(pptx_path).slides:
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if shape.has_table:
                parts.extend(cell.text for row in shape.table.rows for cell in row.cells)
        texts.append("\n".join(parts))
    return texts


def main(pptx_path):
    mapping = json.loads((ROOT / "slides/values_map.json").read_text())["values"]
    texts = slide_texts(pptx_path)
    cache, failures, checked = {}, [], 0
    for entry in mapping:
        value = expected(entry, cache)
        if entry["slide"] is None:
            continue
        slides = [t for t in texts if entry["slide"] in t]
        if not slides:
            failures.append(f"{entry['id']}: no slide titled '{entry['slide']}'")
            continue
        checked += 1
        if not any(value in t for t in slides):
            failures.append(f"{entry['id']}: '{value}' (from {entry['source']}:{entry['path']}) not on '{entry['slide']}'")
    for line in failures:
        print("MISMATCH", line)
    print(f"{checked} mapped values checked, {len(failures)} problems")
    return 1 if failures else 0


if __name__ == "__main__":
    deck = sys.argv[1] if len(sys.argv) > 1 else ROOT / "slides/Knowledge_Graphs_Skills_Neuroimaging_v4.pptx"
    sys.exit(main(deck))
